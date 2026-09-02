import json
from pathlib import Path
from typing import Any
import pandas as pd
from datetime import datetime, timezone


class ReportService:
    def __init__(self):
        self.reports_dir = Path("reports")
        self.reports_dir.mkdir(exist_ok=True)

    def generate_report(
        self, title: str, format_type: str, dataset_path: str | None, profile_json: dict
    ) -> str:
        timestamp = datetime.now(timezone.utc).strftime("%Y%m%d%H%M%S")
        safe_title = "".join([c if c.isalnum() else "_" for c in title])
        base_name = f"{safe_title}_{timestamp}"

        df = None
        if dataset_path and Path(dataset_path).exists():
            sfx = Path(dataset_path).suffix.lower()
            if sfx == ".csv":
                df = pd.read_csv(dataset_path)
            elif sfx in (".xls", ".xlsx"):
                df = pd.read_excel(dataset_path)
            elif sfx == ".json":
                df = pd.read_json(dataset_path)

        fmt = format_type.lower()

        if fmt == "csv":
            return self._generate_csv(df, base_name)
        elif fmt in ("excel", "xlsx"):
            return self._generate_excel(df, base_name, profile_json)
        elif fmt == "docx":
            return self._generate_docx(df, base_name, title, profile_json)
        elif fmt == "pdf":
            return self._generate_pdf(df, base_name, title, profile_json)
        elif fmt == "html":
            return self._generate_html(df, base_name, title, profile_json)
        elif fmt == "json":
            return self._generate_json(df, base_name, title, profile_json)
        else:
            raise ValueError(f"Unsupported format: {format_type}")

    def _generate_csv(self, df: pd.DataFrame | None, base_name: str) -> str:
        if df is None:
            raise ValueError("Cannot generate CSV without a valid dataset")
        output_path = self.reports_dir / f"{base_name}.csv"
        df.to_csv(output_path, index=False)
        return str(output_path)

    def _generate_excel(self, df: pd.DataFrame | None, base_name: str, profile_json: dict) -> str:
        if df is None:
            raise ValueError("Cannot generate Excel without a valid dataset")
        output_path = self.reports_dir / f"{base_name}.xlsx"
        with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name="Data", index=False)
            # Add statistics sheet
            stats = profile_json.get("statistics", {})
            if stats:
                rows = []
                for col, vals in stats.items():
                    row = {"Column": col}
                    row.update({k: v for k, v in vals.items() if v != ""})
                    rows.append(row)
                if rows:
                    pd.DataFrame(rows).to_excel(writer, sheet_name="Statistics", index=False)
            # Missing values sheet
            mv = profile_json.get("missing_values", {})
            if mv:
                pd.DataFrame(list(mv.items()), columns=["Column", "Missing Count"]).to_excel(
                    writer, sheet_name="Missing Values", index=False
                )
            # Outliers sheet
            outliers = profile_json.get("outliers", {})
            if outliers:
                pd.DataFrame(list(outliers.items()), columns=["Column", "Outlier Count"]).to_excel(
                    writer, sheet_name="Outliers", index=False
                )
        return str(output_path)

    def _generate_docx(self, df: pd.DataFrame | None, base_name: str, title: str, profile_json: dict) -> str:
        from docx import Document
        output_path = self.reports_dir / f"{base_name}.docx"
        doc = Document()
        doc.add_heading(title, 0)
        doc.add_paragraph(
            f"Report generated on: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')}"
        )
        # Dataset overview
        if df is not None:
            doc.add_heading("Dataset Overview", level=1)
            doc.add_paragraph(f"Rows: {len(df)}")
            doc.add_paragraph(f"Columns: {len(df.columns)}")
            doc.add_paragraph(f"Duplicate rows: {profile_json.get('duplicates', 0)}")
            doc.add_paragraph(f"Column names: {', '.join(df.columns.tolist())}")
        # Missing values
        mv = profile_json.get("missing_values", {})
        if mv:
            doc.add_heading("Missing Values", level=1)
            for col, count in mv.items():
                doc.add_paragraph(f"  {col}: {count} missing")
        # Outliers
        outliers = profile_json.get("outliers", {})
        if outliers and any(v > 0 for v in outliers.values()):
            doc.add_heading("Outlier Analysis", level=1)
            for col, count in outliers.items():
                if count > 0:
                    doc.add_paragraph(f"  {col}: {count} outliers detected")
        # Statistics
        stats = profile_json.get("statistics", {})
        if stats:
            doc.add_heading("Statistical Summary", level=1)
            for col, vals in list(stats.items())[:15]:
                doc.add_heading(f"Column: {col}", level=2)
                for k, v in vals.items():
                    if v != "":
                        doc.add_paragraph(f"  {k}: {v}")
        # Data Quality
        quality = self._compute_data_quality(profile_json, df)
        if quality:
            doc.add_heading("Data Quality Score", level=1)
            doc.add_paragraph(f"Overall Score: {quality['overall_score']}/100")
            for dim, score in quality.get("dimensions", {}).items():
                doc.add_paragraph(f"  {dim}: {score}/100")

        doc.save(output_path)
        return str(output_path)

    def _generate_pdf(self, df: pd.DataFrame | None, base_name: str, title: str, profile_json: dict) -> str:
        from reportlab.lib.pagesizes import letter
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.units import inch
        from reportlab.lib import colors
        from reportlab.platypus import (
            SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
        )

        output_path = self.reports_dir / f"{base_name}.pdf"
        doc_pdf = SimpleDocTemplate(str(output_path), pagesize=letter)
        styles = getSampleStyleSheet()
        story = []

        # Title
        story.append(Paragraph(title, styles["Title"]))
        story.append(Paragraph(
            f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}",
            styles["Normal"]
        ))
        story.append(Spacer(1, 0.3 * inch))

        # Dataset overview
        if df is not None:
            story.append(Paragraph("Dataset Overview", styles["Heading1"]))
            overview_data = [
                ["Metric", "Value"],
                ["Total Rows", str(len(df))],
                ["Total Columns", str(len(df.columns))],
                ["Duplicate Rows", str(profile_json.get("duplicates", 0))],
                ["Columns", ", ".join(df.columns.tolist()[:10]) + ("..." if len(df.columns) > 10 else "")],
            ]
            t = Table(overview_data, colWidths=[2.5 * inch, 4 * inch])
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2563EB")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F5F9")]),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
            ]))
            story.append(t)
            story.append(Spacer(1, 0.2 * inch))

        # Missing values
        mv = profile_json.get("missing_values", {})
        if any(v > 0 for v in mv.values()):
            story.append(Paragraph("Missing Values", styles["Heading1"]))
            mv_data = [["Column", "Missing Count"]] + [[k, str(v)] for k, v in mv.items() if v > 0]
            t = Table(mv_data, colWidths=[3 * inch, 2 * inch])
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#DC2626")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
            ]))
            story.append(t)
            story.append(Spacer(1, 0.2 * inch))

        # Outliers
        outliers = profile_json.get("outliers", {})
        if outliers and any(v > 0 for v in outliers.values()):
            story.append(Paragraph("Outlier Analysis", styles["Heading1"]))
            ol_data = [["Column", "Outlier Count"]] + [[k, str(v)] for k, v in outliers.items() if v > 0]
            t = Table(ol_data, colWidths=[3 * inch, 2 * inch])
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#F59E0B")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
            ]))
            story.append(t)
            story.append(Spacer(1, 0.2 * inch))

        # Statistics
        stats = profile_json.get("statistics", {})
        if stats:
            story.append(Paragraph("Statistical Summary", styles["Heading1"]))
            for col, vals in list(stats.items())[:10]:
                story.append(Paragraph(f"Column: {col}", styles["Heading2"]))
                clean_vals = {k: str(v) for k, v in vals.items() if v != ""}
                if clean_vals:
                    stat_data = [["Stat", "Value"]] + list(clean_vals.items())
                    t = Table(stat_data, colWidths=[2 * inch, 3 * inch])
                    t.setStyle(TableStyle([
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4F46E5")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                        ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ]))
                    story.append(t)
                story.append(Spacer(1, 0.15 * inch))

        doc_pdf.build(story)
        return str(output_path)

    def _generate_html(self, df: pd.DataFrame | None, base_name: str, title: str, profile_json: dict) -> str:
        """Generate a styled HTML report."""
        output_path = self.reports_dir / f"{base_name}.html"
        generated_at = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
        quality = self._compute_data_quality(profile_json, df)

        html_parts = [
            "<!DOCTYPE html>",
            "<html lang='en'><head><meta charset='UTF-8'>",
            f"<title>{title}</title>",
            "<style>",
            "body { font-family: 'Segoe UI', Tahoma, sans-serif; background: #0f172a; color: #e2e8f0; margin: 0; padding: 2rem; }",
            "h1 { color: #60a5fa; border-bottom: 2px solid #1e40af; padding-bottom: 0.5rem; }",
            "h2 { color: #a78bfa; margin-top: 2rem; }",
            "table { border-collapse: collapse; width: 100%; margin: 1rem 0; }",
            "th { background: #1e293b; color: #60a5fa; padding: 0.75rem; text-align: left; border: 1px solid #334155; }",
            "td { padding: 0.5rem 0.75rem; border: 1px solid #334155; }",
            "tr:nth-child(even) { background: #1e293b; }",
            "tr:hover { background: #334155; }",
            ".metric-card { display: inline-block; background: #1e293b; border: 1px solid #334155; border-radius: 12px; padding: 1.5rem; margin: 0.5rem; min-width: 180px; }",
            ".metric-value { font-size: 1.8rem; font-weight: bold; color: #60a5fa; }",
            ".metric-label { font-size: 0.85rem; color: #94a3b8; margin-top: 0.25rem; }",
            ".quality-score { font-size: 2.5rem; font-weight: bold; text-align: center; padding: 1rem; }",
            ".badge { display: inline-block; padding: 0.25rem 0.75rem; border-radius: 9999px; font-size: 0.75rem; font-weight: bold; }",
            ".badge-good { background: #065f46; color: #6ee7b7; }",
            ".badge-warn { background: #78350f; color: #fbbf24; }",
            ".badge-bad { background: #7f1d1d; color: #fca5a5; }",
            "footer { margin-top: 3rem; padding-top: 1rem; border-top: 1px solid #334155; color: #64748b; font-size: 0.85rem; }",
            "</style></head><body>",
            f"<h1>{title}</h1>",
            f"<p style='color:#94a3b8'>Generated: {generated_at}</p>",
        ]

        # Overview metrics
        if df is not None:
            html_parts.append("<div style='margin: 1.5rem 0'>")
            html_parts.append(f"<div class='metric-card'><div class='metric-value'>{len(df):,}</div><div class='metric-label'>Total Rows</div></div>")
            html_parts.append(f"<div class='metric-card'><div class='metric-value'>{len(df.columns)}</div><div class='metric-label'>Columns</div></div>")
            html_parts.append(f"<div class='metric-card'><div class='metric-value'>{profile_json.get('duplicates', 0)}</div><div class='metric-label'>Duplicate Rows</div></div>")
            if quality:
                score = quality["overall_score"]
                color = "#6ee7b7" if score >= 80 else ("#fbbf24" if score >= 60 else "#fca5a5")
                html_parts.append(f"<div class='metric-card'><div class='metric-value' style='color:{color}'>{score}/100</div><div class='metric-label'>Data Quality Score</div></div>")
            html_parts.append("</div>")

        # Missing values
        mv = profile_json.get("missing_values", {})
        has_missing = any(v > 0 for v in mv.values())
        if has_missing:
            html_parts.append("<h2>Missing Values</h2>")
            html_parts.append("<table><thead><tr><th>Column</th><th>Missing Count</th><th>Status</th></tr></thead><tbody>")
            for col, count in mv.items():
                if count > 0:
                    badge = "badge-bad" if count > 10 else "badge-warn"
                    html_parts.append(f"<tr><td>{col}</td><td>{count}</td><td><span class='badge {badge}'>{count} missing</span></td></tr>")
            html_parts.append("</tbody></table>")

        # Statistics
        stats = profile_json.get("statistics", {})
        if stats:
            html_parts.append("<h2>Statistical Summary</h2>")
            html_parts.append("<table><thead><tr><th>Column</th>")
            stat_keys = ["count", "mean", "std", "min", "25%", "50%", "75%", "max"]
            for k in stat_keys:
                html_parts.append(f"<th>{k}</th>")
            html_parts.append("</tr></thead><tbody>")
            for col, vals in list(stats.items())[:20]:
                html_parts.append(f"<tr><td><strong>{col}</strong></td>")
                for k in stat_keys:
                    v = vals.get(k, "")
                    if isinstance(v, float):
                        v = f"{v:,.2f}"
                    html_parts.append(f"<td>{v}</td>")
                html_parts.append("</tr>")
            html_parts.append("</tbody></table>")

        # Data preview
        if df is not None and len(df) > 0:
            html_parts.append("<h2>Data Preview (First 20 Rows)</h2>")
            preview = df.head(20)
            html_parts.append(preview.to_html(classes="", index=False, border=0))

        html_parts.append(f"<footer>Report generated by Unified Enterprise AI Analytics Platform — {generated_at}</footer>")
        html_parts.append("</body></html>")

        output_path.write_text("\n".join(html_parts), encoding="utf-8")
        return str(output_path)

    def _generate_json(self, df: pd.DataFrame | None, base_name: str, title: str, profile_json: dict) -> str:
        """Generate a JSON report with all metadata."""
        output_path = self.reports_dir / f"{base_name}.json"
        quality = self._compute_data_quality(profile_json, df)

        report = {
            "title": title,
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "dataset": {
                "rows": len(df) if df is not None else 0,
                "columns": len(df.columns) if df is not None else 0,
                "column_names": list(df.columns) if df is not None else [],
                "duplicates": profile_json.get("duplicates", 0),
            },
            "profiling": {
                "schema": profile_json.get("schema", {}),
                "statistics": profile_json.get("statistics", {}),
                "missing_values": profile_json.get("missing_values", {}),
                "unique_values": profile_json.get("unique_values", {}),
                "outliers": profile_json.get("outliers", {}),
                "class_imbalance": profile_json.get("class_imbalance", {}),
                "correlations": profile_json.get("correlations", {}),
            },
            "data_quality": quality,
        }

        output_path.write_text(json.dumps(report, indent=2, default=str), encoding="utf-8")
        return str(output_path)

    def _compute_data_quality(self, profile_json: dict, df: pd.DataFrame | None) -> dict[str, Any]:
        """Compute a data quality score from profiling results."""
        if df is None or not profile_json:
            return {}

        total_cells = len(df) * len(df.columns) if len(df.columns) > 0 else 1
        mv = profile_json.get("missing_values", {})
        total_missing = sum(mv.values())
        completeness = max(0, round((1 - total_missing / total_cells) * 100, 1))

        duplicates = profile_json.get("duplicates", 0)
        uniqueness = max(0, round((1 - duplicates / max(len(df), 1)) * 100, 1))

        outliers = profile_json.get("outliers", {})
        total_outliers = sum(outliers.values())
        consistency = max(0, round((1 - total_outliers / max(total_cells, 1)) * 100, 1))

        overall = round((completeness * 0.4 + uniqueness * 0.3 + consistency * 0.3), 1)

        return {
            "overall_score": min(100, overall),
            "dimensions": {
                "completeness": completeness,
                "uniqueness": uniqueness,
                "consistency": consistency,
            },
        }
